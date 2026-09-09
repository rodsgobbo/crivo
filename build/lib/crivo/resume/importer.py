"""Importacao de curriculo a partir do Drive, de Word ou de PDF.

As tres origens sao reduzidas ao mesmo artefato: texto puro mais a procedencia.
Nada adiante precisa saber de onde o documento veio, e trocar de origem nao
alcanca o extrator nem o consolidador.

A ordem das verificacoes e deliberada. Formato, tamanho e frequencia sao
conferidos antes de qualquer extracao de texto, e a extracao de texto acontece
antes da extracao estruturada, que e a etapa que custa chamada de modelo. Cada
recusa acontece o mais cedo possivel.

A recusa por ausencia de camada de texto e uma decisao de produto e nao um
detalhe: um curriculo digitalizado produziria extracao vazia, que viraria perfil
pobre, que geraria buscas ruins que o usuario atribuiria ao sistema em vez de ao
documento.
"""

from __future__ import annotations

import hashlib
import io
import uuid
from dataclasses import dataclass
from datetime import date, datetime, timezone
from typing import Protocol

from ..store.repository import Repository

#: Extensoes aceitas por formato de origem.
EXTENSOES = {"docx": ("docx",), "pdf": ("pdf",)}

ORIGENS = ("gdoc", "docx", "pdf")


class ResumeImportError(Exception):
    """Documento recusado antes de virar registro de curriculo."""


@dataclass(frozen=True)
class ImportedResume:
    """Curriculo aceito e gravado."""

    resume_id: str
    origem: str
    texto: str
    hash_conteudo: str
    importado_em: str


class DriveClient(Protocol):
    """Porta do Drive. Pede apenas o escopo de leitura do documento escolhido."""

    #: Escopo minimo necessario para exportar o documento selecionado.
    scope: str

    def export_text(self, file_id: str, access_token: str) -> str: ...


class GoogleDriveClient:
    """Exporta o documento selecionado como texto puro."""

    #: Escopo por arquivo: alcanca apenas o que o usuario escolheu no seletor,
    #: e nao o conteudo do Drive inteiro.
    scope = "https://www.googleapis.com/auth/drive.file"
    EXPORT_URL = "https://www.googleapis.com/drive/v3/files/{file_id}/export"

    def __init__(self, http=None) -> None:
        self._http = http

    def export_text(self, file_id: str, access_token: str) -> str:
        if self._http is None:  # pragma: no cover - exige rede
            import httpx

            self._http = httpx.Client(timeout=30)
        response = self._http.get(
            self.EXPORT_URL.format(file_id=file_id),
            params={"mimeType": "text/plain"},
            headers={"Authorization": f"Bearer {access_token}"},
        )
        if response.status_code != 200:
            raise ResumeImportError(
                f"o Drive recusou a exportacao do documento: HTTP "
                f"{response.status_code}"
            )
        return response.text


def extract_docx_text(data: bytes) -> str:
    """Le o texto de um arquivo Word, paragrafos e tabelas."""
    try:
        import docx
    except ImportError as exc:  # pragma: no cover - dependencia declarada
        raise ResumeImportError("suporte a Word indisponivel neste ambiente") from exc
    try:
        documento = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ResumeImportError(f"arquivo Word ilegivel: {type(exc).__name__}") from exc
    partes = [p.text for p in documento.paragraphs]
    for tabela in documento.tables:
        for linha in tabela.rows:
            partes.extend(celula.text for celula in linha.cells)
    return "\n".join(parte for parte in partes if parte and parte.strip())


def extract_pdf_text(data: bytes) -> str:
    """Le o texto de um PDF. Documento digitalizado devolve pouco ou nada."""
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependencia declarada
        raise ResumeImportError("suporte a PDF indisponivel neste ambiente") from exc
    try:
        leitor = PdfReader(io.BytesIO(data))
        paginas = [pagina.extract_text() or "" for pagina in leitor.pages]
    except Exception as exc:
        raise ResumeImportError(f"arquivo PDF ilegivel: {type(exc).__name__}") from exc
    return "\n".join(p for p in paginas if p.strip())


class ResumeImporter:
    """Aceita, valida e grava curriculos sob o usuario que os trouxe."""

    def __init__(self, connection, config, drive: DriveClient | None = None) -> None:
        self._repository = Repository(connection)
        self._config = config.resume
        self._drive = drive or GoogleDriveClient()

    # ------------------------------------------------------------- origens
    def import_from_drive(
        self, user_id: str, file_id: str, access_token: str
    ) -> ImportedResume:
        self._check_frequency(user_id)
        self._check_origin("gdoc")
        texto = self._drive.export_text(file_id, access_token)
        return self._accept(user_id, "gdoc", texto)

    def import_upload(
        self, user_id: str, filename: str, data: bytes
    ) -> ImportedResume:
        self._check_frequency(user_id)
        origem = self._detect_origin(filename)
        self._check_size(data)
        texto = (
            extract_docx_text(data) if origem == "docx" else extract_pdf_text(data)
        )
        return self._accept(user_id, origem, texto)

    @property
    def required_drive_scope(self) -> str:
        return self._drive.scope

    # ------------------------------------------------------------ validacao
    def _detect_origin(self, filename: str) -> str:
        extensao = (filename or "").rsplit(".", 1)[-1].lower()
        for origem, extensoes in EXTENSOES.items():
            if extensao in extensoes:
                return origem
        raise ResumeImportError(
            f"formato {extensao or 'desconhecido'!r} nao aceito; "
            f"formatos aceitos: {', '.join(self._config.formatos_aceitos)}"
        )

    def _check_origin(self, origem: str) -> None:
        if origem not in self._config.formatos_aceitos:
            raise ResumeImportError(
                f"origem {origem!r} nao aceita; "
                f"formatos aceitos: {', '.join(self._config.formatos_aceitos)}"
            )

    def _check_size(self, data: bytes) -> None:
        limite = self._config.tamanho_maximo_bytes
        if len(data) > limite:
            raise ResumeImportError(
                f"arquivo com {len(data)} bytes excede o limite de {limite} bytes"
            )

    def _check_frequency(self, user_id: str) -> None:
        # Data em UTC, nao local: todo carimbo gravado no banco e UTC, e contar
        # "hoje" no fuso local fazia o contador zerar antes da virada do dia --
        # no horario de Brasilia, todo limite diario deixava de valer das 21h a
        # meia-noite.
        hoje = datetime.now(timezone.utc).date().isoformat()
        usadas = len(
            self._repository.for_user(user_id).select(
                "resumes", where="substr(importado_em, 1, 10) = ?", params=(hoje,)
            )
        )
        limite = self._config.importacoes_por_dia
        if usadas >= limite:
            raise ResumeImportError(
                f"limite de {limite} importacoes por dia atingido; "
                "novas importacoes serao aceitas a partir de amanha"
            )

    def _check_text_layer(self, texto: str) -> None:
        minimo = self._config.comprimento_minimo_texto
        if len(texto.strip()) < minimo:
            raise ResumeImportError(
                f"o documento rendeu {len(texto.strip())} caracteres, abaixo do "
                f"minimo de {minimo}; provavelmente e uma imagem digitalizada sem "
                "camada de texto. Envie um arquivo com texto selecionavel"
            )

    # -------------------------------------------------------------- gravacao
    def _accept(self, user_id: str, origem: str, texto: str) -> ImportedResume:
        self._check_text_layer(texto)
        resume_id = str(uuid.uuid4())
        digest = content_hash(texto)
        importado_em = datetime.now(timezone.utc).isoformat(timespec="seconds")
        self._repository.for_user(user_id).insert(
            "resumes",
            {
                "resume_id": resume_id,
                "texto": texto,
                "origem": origem,
                "hash_conteudo": digest,
                "importado_em": importado_em,
            },
        )
        return ImportedResume(
            resume_id=resume_id,
            origem=origem,
            texto=texto,
            hash_conteudo=digest,
            importado_em=importado_em,
        )


def content_hash(texto: str) -> str:
    """Identidade do conteudo, usada para reusar extracao ja feita."""
    normalizado = " ".join((texto or "").split())
    return hashlib.sha256(normalizado.encode("utf-8")).hexdigest()
